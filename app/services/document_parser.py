import os
import tempfile
from typing import Optional


async def extract_text_from_file(file_content: bytes, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".txt":
        return file_content.decode("utf-8", errors="ignore")

    if ext == ".md":
        return file_content.decode("utf-8", errors="ignore")

    if ext == ".docx":
        return _extract_from_docx(file_content)

    if ext == ".pdf":
        return _extract_from_pdf(file_content)

    return file_content.decode("utf-8", errors="ignore")


def _extract_from_docx(file_content: bytes) -> str:
    from docx import Document
    import io

    doc = Document(io.BytesIO(file_content))
    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    return "\n".join(paragraphs)


def _extract_from_pdf(file_content: bytes) -> str:
    from PyPDF2 import PdfReader
    import io

    reader = PdfReader(io.BytesIO(file_content))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text and text.strip():
            pages.append(text.strip())

    return "\n".join(pages)
